"""勤務表の Excel(.xlsx) / 印刷用 PDF 出力。どちらも A4 横 1 ページに収まる体裁。"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from scheduler import COUNTED_CODES, WEEKDAY_JA, ScheduleInput, roster_summary

SUMMARY_COLS = ["公休", "PB", "早番", "ハ番", "半休"]


def _hex(color: str, default: str = "FFFFFF") -> str:
    c = (color or "").lstrip("#").upper()
    return c if len(c) == 6 else default


def _day_color(di) -> str | None:
    if di.is_holiday or di.weekday == 6:
        return "F4B6B6"
    if di.weekday == 5:
        return "BDD7EE"
    return None


def export_excel(inp: ScheduleInput, roster: dict[int, dict[int, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{inp.year}年{inp.month}月"
    n = inp.n_days
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    font = "Meiryo UI"
    last_col = 1 + n + len(SUMMARY_COLS)

    ws.cell(1, 1, f"放射線科 勤務表  {inp.year}年{inp.month}月").font = Font(name=font, size=16, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)

    # ヘッダ(日付・曜日)
    ws.cell(3, 1, "氏名")
    ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    for di in inp.calendar:
        col = 1 + di.day
        ws.cell(3, col, di.day)
        ws.cell(4, col, "祝" if di.is_holiday and di.weekday != 6 else WEEKDAY_JA[di.weekday])
        fill = _day_color(di)
        for r in (3, 4):
            if fill:
                ws.cell(r, col).fill = PatternFill("solid", fgColor=fill)
    for i, name in enumerate(SUMMARY_COLS):
        col = 2 + n + i
        ws.cell(3, col, name)
        ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
    for r in (3, 4):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=font, size=9, bold=True)
            cell.alignment = center
            cell.border = border

    # 本体
    summary = {row["技師"]: row for row in roster_summary(inp, roster)}
    row = 5
    for s in inp.staff:
        ws.cell(row, 1, s["name"])
        for di in inp.calendar:
            code = roster.get(s["id"], {}).get(di.day, "")
            cell = ws.cell(row, 1 + di.day, code)
            color = _hex(inp.shifts.get(code, {}).get("color", ""))
            if color != "FFFFFF":
                cell.fill = PatternFill("solid", fgColor=color)
        sm = summary[s["name"]]
        for i, key in enumerate(SUMMARY_COLS):
            ws.cell(row, 2 + n + i, sm[key])
        for c in range(1, last_col + 1):
            cell = ws.cell(row, c)
            cell.font = Font(name=font, size=10)
            cell.alignment = center
            cell.border = border
        row += 1

    # 日別集計
    row += 0
    for code in COUNTED_CODES:
        ws.cell(row, 1, f"{inp.name_of(code)}計")
        for di in inp.calendar:
            ws.cell(row, 1 + di.day, sum(1 for s in inp.staff if roster.get(s["id"], {}).get(di.day) == code))
        for c in range(1, last_col + 1):
            cell = ws.cell(row, c)
            cell.font = Font(name=font, size=8, color="404040")
            cell.alignment = center
            cell.border = border
        row += 1

    # 凡例
    row += 1
    legend = "  ".join(f"[{c}]{v['name']}" + (f"({v['start_time']}~)" if v.get("start_time") else "")
                       for c, v in inp.shifts.items())
    ws.cell(row, 1, "凡例: " + legend).font = Font(name=font, size=8)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)

    # 列幅・行高
    ws.column_dimensions["A"].width = 12
    for d in range(1, n + 1):
        ws.column_dimensions[get_column_letter(1 + d)].width = 4.2
    for i in range(len(SUMMARY_COLS)):
        ws.column_dimensions[get_column_letter(2 + n + i)].width = 5.5
    ws.row_dimensions[1].height = 26
    for r in range(5, 5 + len(inp.staff)):
        ws.row_dimensions[r].height = 22

    # 印刷設定: A4 横・1ページに収める
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.freeze_panes = "B5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(inp: ScheduleInput, roster: dict[int, dict[int, str]]) -> bytes:
    """印刷用 PDF。日本語は ReportLab 内蔵の CID フォント(外部フォント不要)を使用。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    fname = "HeiseiKakuGo-W5"
    n = inp.n_days
    summary = {row["技師"]: row for row in roster_summary(inp, roster)}

    head1 = ["氏名"] + [str(di.day) for di in inp.calendar] + SUMMARY_COLS
    head2 = [""] + ["祝" if di.is_holiday and di.weekday != 6 else WEEKDAY_JA[di.weekday] for di in inp.calendar] + [""] * len(SUMMARY_COLS)
    data = [head1, head2]
    for s in inp.staff:
        row = [s["name"]] + [roster.get(s["id"], {}).get(d, "") for d in range(1, n + 1)]
        row += [summary[s["name"]][k] for k in SUMMARY_COLS]
        data.append(row)
    for code in COUNTED_CODES:
        data.append([f"{inp.name_of(code)}計"] + [
            sum(1 for s in inp.staff if roster.get(s["id"], {}).get(d) == code) for d in range(1, n + 1)
        ] + [""] * len(SUMMARY_COLS))

    page_w = landscape(A4)[0] - 20 * mm
    name_w, sum_w = 22 * mm, 9 * mm
    day_w = (page_w - name_w - sum_w * len(SUMMARY_COLS)) / n
    col_w = [name_w] + [day_w] * n + [sum_w] * len(SUMMARY_COLS)

    def rl(hexcolor: str):
        return colors.HexColor("#" + _hex(hexcolor))

    style = [
        ("FONT", (0, 0), (-1, -1), fname, 7),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#EEEEEE")),
        ("SPAN", (0, 0), (0, 1)),
    ] + [("SPAN", (1 + n + i, 0), (1 + n + i, 1)) for i in range(len(SUMMARY_COLS))]
    for di in inp.calendar:
        c = _day_color(di)
        if c:
            style.append(("BACKGROUND", (di.day, 0), (di.day, 1), rl(c)))
    n_staff = len(inp.staff)
    for i, s in enumerate(inp.staff):
        for d in range(1, n + 1):
            code = roster.get(s["id"], {}).get(d, "")
            color = _hex(inp.shifts.get(code, {}).get("color", ""))
            if color != "FFFFFF":
                style.append(("BACKGROUND", (d, 2 + i), (d, 2 + i), rl(color)))
    style += [("FONTSIZE", (0, 2 + n_staff), (-1, -1), 6), ("TEXTCOLOR", (0, 2 + n_staff), (-1, -1), colors.HexColor("#404040"))]

    row_h = min(9 * mm, (landscape(A4)[1] - 55 * mm) / len(data))
    table = Table(data, colWidths=col_w, rowHeights=row_h, repeatRows=2)
    table.setStyle(TableStyle(style))

    title_style = ParagraphStyle("t", fontName=fname, fontSize=15, leading=20)
    legend_style = ParagraphStyle("l", fontName=fname, fontSize=7, leading=10)
    legend = "　".join(f"[{c}]{v['name']}" + (f"({v['start_time']}~)" if v.get("start_time") else "")
                      for c, v in inp.shifts.items())
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f"勤務表 {inp.year}年{inp.month}月")
    doc.build([Paragraph(f"放射線科 勤務表　{inp.year}年{inp.month}月", title_style), table,
               Paragraph("凡例: " + legend, legend_style)])
    return buf.getvalue()
